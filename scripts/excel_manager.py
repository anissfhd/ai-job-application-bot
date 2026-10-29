"""
Gestion du fichier Excel de suivi des candidatures (input/entreprises.xlsx).

Principes :
- Le fichier est RELU à chaque opération : tu peux ajouter des offres ou
  modifier un statut à la main pendant que le robot tourne.
- Les colonnes sont retrouvées par leur NOM (tu peux les réordonner).
- Si le fichier est ouvert dans Excel (verrouillé sous Windows), l'écriture
  est retentée plusieurs fois au lieu de planter.

Utilisation en ligne de commande :
    python scripts/excel_manager.py --init     # crée le modèle vide
    python scripts/excel_manager.py --stats    # affiche un résumé des statuts
"""

from __future__ import annotations

import argparse
import logging
import os
import time
import unicodedata
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


def enregistrer_sans_risque(wb, chemin: Path, sauvegardes_max: int = 72) -> None:
    """
    Enregistre le classeur SANS jamais pouvoir l'abîmer.

    `wb.save(chemin)` réécrit le fichier en place : un Ctrl+C (ou une coupure) en plein
    enregistrement laissait un fichier à moitié écrit, illisible — tout l'Excel était perdu.
    On écrit donc d'abord un fichier temporaire complet, puis on le substitue d'un coup à
    l'original (os.replace est atomique) : l'Excel est toujours soit l'ancien, soit le nouveau.

    Au passage, une copie de l'ancien est gardée une fois par heure dans « sauvegardes/ »
    (les `sauvegardes_max` plus récentes), pour pouvoir revenir en arrière.
    PermissionError (fichier ouvert dans Excel) remonte telle quelle : l'appelant réessaie.
    """
    chemin = Path(chemin)
    tmp = chemin.with_name(chemin.stem + ".enregistrement.xlsx")
    try:
        wb.save(tmp)
        _sauvegarde_horaire(chemin, sauvegardes_max)
        os.replace(tmp, chemin)
    except BaseException:
        try:
            tmp.unlink(missing_ok=True)
        except OSError:
            pass
        raise


def _sauvegarde_horaire(chemin: Path, garder: int) -> None:
    """Copie l'Excel actuel (s'il est lisible) dans sauvegardes/, une fois par heure."""
    import shutil
    import zipfile
    if not chemin.exists():
        return
    dossier = chemin.parent / "sauvegardes"
    cible = dossier / f"{chemin.stem}_{datetime.now():%Y-%m-%d_%Hh}.xlsx"
    if cible.exists():
        return
    try:
        with zipfile.ZipFile(chemin) as z:
            if "[Content_Types].xml" not in z.namelist():
                return               # fichier déjà abîmé : on ne l'archive pas
    except (zipfile.BadZipFile, OSError):
        return
    dossier.mkdir(exist_ok=True)
    shutil.copy2(chemin, cible)
    for vieux in sorted(dossier.glob(f"{chemin.stem}_*.xlsx"))[:-garder]:
        try:
            vieux.unlink()
        except OSError:
            pass


# Colonnes demandées + 2 colonnes techniques (nombre d'essais, détail de l'erreur)
COLONNES = [
    "Nom_Entreprise", "Lien_Offre", "Email", "Téléphone",
    "Description_Entreprise", "Description_Offre",
    "CV_PDF_Path", "LM_PDF_Path", "Statut", "Date_Traitement",
    "Tentatives", "Commentaire",
]

# Statuts possibles
A_TRAITER = "À traiter"          # (ou cellule vide) -> sera traité
EN_COURS = "En cours"
POSTULE = "Postulé"
PRET_TEST = "Prêt (test)"        # formulaire rempli mais non envoyé (soumission_automatique: false)
ERREUR = "Erreur"                # sera retenté jusqu'à max_tentatives_par_offre
ACTION_MANUELLE = "Action manuelle"  # CAPTCHA, connexion requise, info manquante…
OFFRE_EXPIREE = "Offre expirée"
IGNORE = "Ignoré"                # mets ce statut à la main pour sauter une ligne
FAIT = "Fait"                    # tu considères l'offre réglée : le robot n'y touche plus
A_VERIFIER = "À vérifier"          # offre scrapée en attente du tri de domaine (verify_eligibility.py) : jamais prise

STATUTS = [A_TRAITER, POSTULE, FAIT, IGNORE, ERREUR, ACTION_MANUELLE, PRET_TEST, EN_COURS, OFFRE_EXPIREE,
           A_VERIFIER]


def statut_normalise(valeur) -> str:
    """
    Statut tel que le robot doit le comprendre, quelle que soit l'écriture dans l'Excel :
    « à traiter », « A traiter », « À TRAITER », « à  traiter » -> « À traiter ». Sans ça, une
    ligne écrite à la main avec une autre casse ou sans accent était sautée sans aucun message.
    Un statut inconnu est renvoyé tel quel.
    """
    brut = str(valeur or "").strip()
    cle = " ".join(unicodedata.normalize("NFKD", brut).encode("ascii", "ignore").decode().lower().split())
    for statut in STATUTS:
        if cle == " ".join(unicodedata.normalize("NFKD", statut).encode("ascii", "ignore").decode().lower().split()):
            return statut
    return brut

LARGEURS = {
    "Nom_Entreprise": 25, "Lien_Offre": 45, "Email": 28, "Téléphone": 16,
    "Description_Entreprise": 50, "Description_Offre": 60,
    "CV_PDF_Path": 40, "LM_PDF_Path": 40, "Statut": 16, "Date_Traitement": 17,
    "Tentatives": 11, "Commentaire": 60,
}

EXEMPLES = [
    ["Exemple SAS (à supprimer)", "https://www.welcometothejungle.com/fr/companies/exemple/jobs/stage-pfe",
     "", "", "", "", "", "", IGNORE, "", 0, "Ligne d'exemple : remplace-la par tes offres"],
]


# VERROU COMMUN sur l'Excel : le robot ET le scraper (input/Scraping) lisent tout le classeur,
# le modifient et le réenregistrent en entier. S'ils le font en même temps, le dernier à
# enregistrer efface ce que l'autre vient d'écrire (une offre ajoutée, un statut « Postulé »).
# Protocole partagé, À GARDER IDENTIQUE des deux côtés : un fichier « <excel>.verrou » créé de
# façon atomique ; un verrou plus vieux que VERROU_PERIME secondes est celui d'un programme
# arrêté brutalement, on le reprend.
VERROU_PERIME = 600
JOURNAL = "job_automation"


def _processus_vivant(pid: int) -> bool:
    """Le programme qui a posé le verrou tourne-t-il encore ? (Dans le doute : oui.)"""
    if pid <= 0:
        return True
    if os.name == "nt":
        import ctypes
        k = ctypes.windll.kernel32
        h = k.OpenProcess(0x1000, False, pid)          # PROCESS_QUERY_LIMITED_INFORMATION
        if not h:
            return k.GetLastError() == 5               # accès refusé = il existe
        code = ctypes.c_ulong()
        ok = k.GetExitCodeProcess(h, ctypes.byref(code))
        k.CloseHandle(h)
        return not ok or code.value == 259             # 259 = STILL_ACTIVE
    try:
        os.kill(pid, 0)
        return True
    except OSError:
        return False


def _verrou_orphelin(fichier: Path) -> bool:
    """Verrou dont le programme n'existe plus (fenêtre fermée, Ctrl+C, plantage) ou trop vieux."""
    try:
        if time.time() - fichier.stat().st_mtime > VERROU_PERIME:
            return True
        morceaux = fichier.read_text(encoding="utf-8", errors="replace").split()
        return len(morceaux) >= 2 and morceaux[1].isdigit() and not _processus_vivant(int(morceaux[1]))
    except OSError:
        return False


@contextmanager
def verrou_excel(chemin: Path, attente_max: float = 900.0):
    fichier = Path(str(chemin) + ".verrou")
    debut = time.monotonic()
    signale = False
    while True:
        try:
            fd = os.open(str(fichier), os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.write(fd, f"robot {os.getpid()} {datetime.now():%Y-%m-%d %H:%M:%S}".encode())
            os.close(fd)
            break
        except FileExistsError:
            if _verrou_orphelin(fichier):
                try:
                    fichier.unlink(missing_ok=True)
                except OSError:
                    pass
                continue
            if not signale and time.monotonic() - debut > 5:
                signale = True
                try:
                    qui = fichier.read_text(encoding="utf-8", errors="replace").strip()
                except OSError:
                    qui = "?"
                logging.getLogger(JOURNAL).warning(
                    f"Excel verrouillé par un autre programme ({qui}) — attente de sa fin…")
            if time.monotonic() - debut > attente_max:
                raise ExcelVerrouille(f"Verrou {fichier.name} toujours pris après {attente_max:.0f} s.")
            time.sleep(0.3)
    try:
        yield
    finally:
        try:
            fichier.unlink(missing_ok=True)
        except OSError:
            pass


class ExcelVerrouille(Exception):
    """Le fichier Excel est ouvert ailleurs et n'a pas pu être enregistré."""


def creer_modele(chemin: Path, avec_exemple: bool = True) -> None:
    """Crée un fichier Excel vide et bien formaté avec toutes les colonnes."""
    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Offres"
    ws.append(COLONNES)

    entete_fill = PatternFill("solid", fgColor="1F4E78")
    for i, nom in enumerate(COLONNES, start=1):
        cellule = ws.cell(row=1, column=i)
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = entete_fill
        cellule.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cellule.column_letter].width = LARGEURS.get(nom, 20)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLONNES)).column_letter}1"

    # Liste déroulante sur la colonne Statut
    col_statut = ws.cell(row=1, column=COLONNES.index("Statut") + 1).column_letter
    validation = DataValidation(type="list", formula1='"' + ",".join(STATUTS) + '"', allow_blank=True)
    validation.add(f"{col_statut}2:{col_statut}5000")
    ws.add_data_validation(validation)

    if avec_exemple:
        for ligne in EXEMPLES:
            ws.append(ligne)

    # Onglet d'aide
    aide = wb.create_sheet("Aide")
    aide.column_dimensions["A"].width = 22
    aide.column_dimensions["B"].width = 100
    aide.append(["Statut", "Signification"])
    for statut, texte in [
        ("(vide) / À traiter", "L'offre sera traitée par le robot."),
        (EN_COURS, "Le robot travaille dessus (remis à « À traiter » au redémarrage s'il a planté)."),
        (POSTULE, "Candidature envoyée."),
        (PRET_TEST, "Formulaire rempli mais pas envoyé (mode test). Remets « À traiter » pour envoyer."),
        (ERREUR, "Échec technique ; retenté automatiquement jusqu'au nombre max de tentatives."),
        (ACTION_MANUELLE, "Le robot ne peut pas finir seul (CAPTCHA, compte à créer, info manquante…). Voir Commentaire."),
        (OFFRE_EXPIREE, "L'offre n'est plus en ligne."),
        (IGNORE, "Ligne ignorée par le robot."),
    ]:
        aide.append([statut, texte])
    aide.append([])
    aide.append(["Astuce", "Seules les colonnes Nom_Entreprise et Lien_Offre sont obligatoires. "
                           "Si tu remplis toi-même les descriptions, le robot ne les ré-extrait pas."])
    aide["A1"].font = aide["B1"].font = Font(bold=True)

    enregistrer_sans_risque(wb, chemin)


class ExcelManager:
    """Lecture / écriture du fichier de suivi, ligne par ligne."""

    def __init__(self, chemin: Path, logger: logging.Logger | None = None,
                 essais_ecriture: int = 20, pause_ecriture: float = 15.0):
        self.chemin = Path(chemin)
        self.logger = logger or logging.getLogger("job_automation")
        self.essais_ecriture = essais_ecriture
        self.pause_ecriture = pause_ecriture
        if not self.chemin.exists():
            creer_modele(self.chemin)
            self.logger.info(f"Fichier Excel modèle créé : {self.chemin}")
        self._profondeur_verrou = 0
        # Ligne choisie -> son lien (ou email). Le scraper (input/Scraping) peut INSÉRER des
        # offres au-dessus d'une ligne entre deux écritures du robot : on la retrouve alors par
        # ce lien au lieu d'écrire dans la mauvaise ligne.
        self._cles_lignes: dict[int, str] = {}

    @contextmanager
    def _verrou(self):
        """Verrou commun avec le scraper ; réentrant (une méthode verrouillée peut en appeler une autre)."""
        if self._profondeur_verrou:
            self._profondeur_verrou += 1
            try:
                yield
            finally:
                self._profondeur_verrou -= 1
            return
        with verrou_excel(self.chemin):
            self._profondeur_verrou = 1
            try:
                yield
            finally:
                self._profondeur_verrou = 0

    @staticmethod
    def _cle(lien, email) -> str:
        return (str(lien or "").strip() or str(email or "").strip()).lower()

    def _ligne_reelle(self, ws, entetes: dict, ligne: int) -> int:
        """Numéro ACTUEL de la ligne choisie (elle a pu descendre si des offres ont été insérées)."""
        cle = self._cles_lignes.get(ligne)
        if not cle:
            return ligne
        def cle_de(n: int) -> str:
            return self._cle(ws.cell(row=n, column=entetes["Lien_Offre"]).value,
                             ws.cell(row=n, column=entetes["Email"]).value)
        if cle_de(ligne) == cle:
            return ligne
        for n in range(2, ws.max_row + 1):
            if cle_de(n) == cle:
                self.logger.info(f"Excel : la ligne {ligne} a été déplacée en {n} (offres insérées "
                                 f"au-dessus par le scraper) -> écriture au bon endroit.")
                return n
        return ligne

    # -- accès bas niveau -----------------------------------------------------

    def _ouvrir(self):
        """Ouvre le classeur et renvoie (classeur, feuille, {nom_colonne: index})."""
        wb = load_workbook(self.chemin)
        ws = wb["Offres"] if "Offres" in wb.sheetnames else wb.worksheets[0]
        entetes = {}
        for cellule in ws[1]:
            if cellule.value:
                entetes[str(cellule.value).strip()] = cellule.column
        # Ajoute automatiquement les colonnes manquantes à droite
        for nom in COLONNES:
            if nom not in entetes:
                col = ws.max_column + 1
                ws.cell(row=1, column=col, value=nom).font = Font(bold=True)
                entetes[nom] = col
        return wb, ws, entetes

    def _enregistrer(self, wb) -> None:
        """Enregistre en réessayant si le fichier est verrouillé (ouvert dans Excel)."""
        for essai in range(1, self.essais_ecriture + 1):
            try:
                enregistrer_sans_risque(wb, self.chemin)
                return
            except PermissionError:
                self.logger.warning(
                    f"Excel verrouillé (ouvert dans Excel ?) — essai {essai}/{self.essais_ecriture}, "
                    f"nouvelle tentative dans {self.pause_ecriture:.0f} s. Ferme le fichier pour débloquer."
                )
                time.sleep(self.pause_ecriture)
        raise ExcelVerrouille(f"Impossible d'enregistrer {self.chemin} (fichier verrouillé).")

    @staticmethod
    def _valeur(ws, ligne: int, entetes: dict, nom: str):
        valeur = ws.cell(row=ligne, column=entetes[nom]).value
        return "" if valeur is None else valeur

    def _ligne_en_dict(self, ws, ligne: int, entetes: dict) -> dict:
        donnees = {nom: self._valeur(ws, ligne, entetes, nom) for nom in COLONNES}
        donnees["_ligne"] = ligne
        return donnees

    # -- API publique ----------------------------------------------------------

    def lire_toutes(self) -> list[dict]:
        """Renvoie toutes les lignes non vides sous forme de dictionnaires."""
        with self._verrou():
            return self._lire_toutes()

    def _lire_toutes(self) -> list[dict]:
        wb, ws, entetes = self._ouvrir()
        lignes = []
        for ligne in range(2, ws.max_row + 1):
            d = self._ligne_en_dict(ws, ligne, entetes)
            if str(d["Nom_Entreprise"]).strip() or str(d["Lien_Offre"]).strip() or str(d["Email"]).strip():
                lignes.append(d)
        wb.close()
        return lignes

    def prochaine_ligne(self, max_tentatives: int = 2, exclure=None) -> dict | None:
        """
        Renvoie la prochaine ligne à traiter (il faut au moins un lien d'offre ou un email) :
        1. d'abord les lignes au statut vide ou « À traiter » ;
        2. sinon, les lignes « Erreur » ayant eu moins de `max_tentatives` essais.
        `exclure(ligne) -> bool` : lignes sautées pour l'instant (ex. offres Easy Apply quand la
        limite quotidienne de LinkedIn est atteinte). L'ordre de l'Excel reste celui des autres.
        """
        a_retenter = None
        for d in self.lire_toutes():
            statut = statut_normalise(d["Statut"])
            if not (str(d["Lien_Offre"]).strip() or str(d["Email"]).strip()):
                continue
            if exclure is not None and exclure(d):
                continue
            if statut in ("", A_TRAITER):
                self._cles_lignes[d["_ligne"]] = self._cle(d["Lien_Offre"], d["Email"])
                return d
            tentatives = int(d["Tentatives"]) if str(d["Tentatives"]).strip().isdigit() else 0
            if statut == ERREUR and tentatives < max_tentatives and a_retenter is None:
                a_retenter = d
        if a_retenter is not None:
            self._cles_lignes[a_retenter["_ligne"]] = self._cle(a_retenter["Lien_Offre"], a_retenter["Email"])
        return a_retenter

    def mettre_a_jour(self, ligne: int, **valeurs) -> None:
        """
        Met à jour une ligne. Exemple :
            excel.mettre_a_jour(5, Statut="Postulé", CV_PDF_Path="...")
        Les noms avec accent (Téléphone) passent par un dict : **{"Téléphone": "..."}
        """
        with self._verrou():
            wb, ws, entetes = self._ouvrir()
            ligne = self._ligne_reelle(ws, entetes, ligne)
            for nom, valeur in valeurs.items():
                if nom not in entetes:
                    raise KeyError(f"Colonne inconnue : {nom}")
                if isinstance(valeur, str) and len(valeur) > 32000:  # limite d'une cellule Excel
                    valeur = valeur[:32000] + "…"
                ws.cell(row=ligne, column=entetes[nom], value=valeur)
            self._enregistrer(wb)

    def definir_statut(self, ligne: int, statut: str, commentaire: str | None = None) -> None:
        """Change le statut, horodate la ligne et (optionnellement) écrit un commentaire."""
        valeurs = {"Statut": statut, "Date_Traitement": datetime.now().strftime("%Y-%m-%d %H:%M")}
        if commentaire is not None:
            valeurs["Commentaire"] = commentaire
        self.mettre_a_jour(ligne, **valeurs)

    def marquer_en_cours(self, ligne: int) -> int:
        """Passe la ligne « En cours » et incrémente le compteur de tentatives. Renvoie ce compteur."""
        with self._verrou():
            wb, ws, entetes = self._ouvrir()
            ligne = self._ligne_reelle(ws, entetes, ligne)
            brut = ws.cell(row=ligne, column=entetes["Tentatives"]).value
            tentatives = (int(brut) if str(brut or "").strip().isdigit() else 0) + 1
            ws.cell(row=ligne, column=entetes["Tentatives"], value=tentatives)
            ws.cell(row=ligne, column=entetes["Statut"], value=EN_COURS)
            ws.cell(row=ligne, column=entetes["Date_Traitement"], value=datetime.now().strftime("%Y-%m-%d %H:%M"))
            self._enregistrer(wb)
            return tentatives

    def reinitialiser_en_cours(self) -> int:
        """Au démarrage : remet « À traiter » les lignes restées « En cours » après un plantage."""
        with self._verrou():
            return self._reinitialiser_en_cours()

    def _reinitialiser_en_cours(self) -> int:
        wb, ws, entetes = self._ouvrir()
        n = 0
        for ligne in range(2, ws.max_row + 1):
            if statut_normalise(ws.cell(row=ligne, column=entetes["Statut"]).value) == EN_COURS:
                ws.cell(row=ligne, column=entetes["Statut"], value=A_TRAITER)
                n += 1
        if n:
            self._enregistrer(wb)
        else:
            wb.close()
        return n

    def candidatures_du_jour(self) -> int:
        """Nombre de candidatures envoyées aujourd'hui (pour respecter le quota journalier)."""
        aujourdhui = datetime.now().strftime("%Y-%m-%d")
        return sum(
            1 for d in self.lire_toutes()
            if statut_normalise(d["Statut"]) == POSTULE and str(d["Date_Traitement"]).startswith(aujourdhui)
        )

    def statistiques(self) -> dict:
        """Compte les lignes par statut."""
        stats: dict[str, int] = {}
        for d in self.lire_toutes():
            statut = statut_normalise(d["Statut"]) or A_TRAITER
            stats[statut] = stats.get(statut, 0) + 1
        return stats


if __name__ == "__main__":
    import sys as _sys
    try:
        _sys.stdout.reconfigure(encoding="utf-8")  # accents corrects dans la console
    except Exception:
        pass
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from utils import charger_config, chemin_projet

    parser = argparse.ArgumentParser(description="Gestion du fichier Excel de suivi")
    parser.add_argument("--init", action="store_true", help="crée le modèle Excel (n'écrase jamais un fichier existant)")
    parser.add_argument("--stats", action="store_true", help="affiche le nombre d'offres par statut")
    args = parser.parse_args()

    config, _ = charger_config()
    chemin = chemin_projet(config["chemins"]["excel"])
    if args.init:
        if chemin.exists():
            print(f"Le fichier existe déjà, rien n'a été modifié : {chemin}")
        else:
            creer_modele(chemin)
            print(f"Modèle créé : {chemin}")
    if args.stats or not args.init:
        for statut, n in ExcelManager(chemin).statistiques().items():
            print(f"{statut:<18} {n}")

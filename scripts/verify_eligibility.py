"""
verify_eligibility — robot SÉPARÉ qui trie les offres scrapées selon le domaine recherché.

Lancé automatiquement (fenêtre à part) par input/Scraping/scraping.py après chaque passage
Apify — manuel (LANCER_SCRAPING.bat) ou automatique (robot de candidature). Il ne touche qu'aux
lignes « À vérifier » (statut que le robot de candidature ne prend jamais) :

  Étape 1 — TITRES : les offres sont numérotées 1, 2, 3… et envoyées par lots de 30 à DeepSeek
            (son propre navigateur, port 9226) avec les critères de criteres_domaine.json.
            DeepSeek répond true / false / pending pour CHAQUE numéro ; réponse incomplète ou
            mal numérotée -> redemandée.
  Étape 2 — DESCRIPTIONS (offres « pending », ET offres « false » sur leur seul titre, car un
            titre peut tromper) : description LinkedIn gardée par le scraping
            (offres_scrapees.json), sans rouvrir LinkedIn. Son verdict remplace celui du titre.
  Étape 3 — SITE DE L'ENTREPRISE (encore « pending » et pas Easy Apply) : lien « Postuler sur le
            site de l'entreprise » de la page publique LinkedIn, puis la page de l'entreprise ;
            DeepSeek ne juge QUE le poste concerné, même si la page en montre plusieurs.
            CAPTCHA sur une de ces pages : même logique NopeCHA que le robot de candidature
            (voir passer_captcha) ; page anti-robot non franchie -> « Action manuelle ».

Résultat : true -> « À traiter » (le robot de candidature la prendra) ; false, ou toujours
indéterminée -> déplacée dans la feuille « Écartées » avec la raison (le scraping ne la ramène
plus). Chaque offre est suivie par son LIEN, jamais par son numéro de ligne (les lignes bougent).

S'il s'arrête en route (coupure, DeepSeek saturé), les offres restent « À vérifier » et sont
reprises au lancement suivant. Il ne tourne jamais en double, et se ferme seul à la fin.

    python scripts/verify_eligibility.py        (ou LANCER.bat, entrée 12)
"""
from __future__ import annotations

import copy
import importlib.util
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

RACINE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RACINE / "scripts"))

import excel_manager as em                                   # noqa: E402
from deepseek_formulaire import ClientDeepSeekFormulaire     # noqa: E402
from gemini_chat import ErreurGeminiChat                     # noqa: E402
from navigateur import Navigateur                            # noqa: E402
from objectif import courant as objectif_courant            # noqa: E402
from openpyxl.styles import Font                             # noqa: E402
from utils import chemin_projet, charger_config, configurer_logs, extraire_json  # noqa: E402

_spec = importlib.util.spec_from_file_location("scraping", RACINE / "input" / "Scraping" / "scraping.py")
S = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(S)

PORT = 9226                                    # 5e navigateur : jamais un des 4 du robot
PROFIL = "~/.tappi/job-automation-verif"
LOT_TITRES = 30
LOT_DESCRIPTIONS = 5
ESSAIS = 3                                     # réponses invalides redemandées au plus 2 fois
TOURS_MAX = 3                                  # offres arrivées pendant le tri : reprises
VERROU = RACINE / "input" / "Scraping" / "verify_eligibility.verrou"
VERDICTS = {"true": "true", "vrai": "true", "oui": "true", "yes": "true",
            "false": "false", "faux": "false", "non": "false", "no": "false",
            "pending": "pending", "inconnu": "pending", "unknown": "pending"}
RAISON_INCONNU = "domaine non déterminé (titre, description et site de l'entreprise insuffisants)"
# CAPTCHA : mêmes valeurs que le robot de candidature (candidature.py).
DELAI_CASE_VALIDEE = 20.0      # après le clic sur la case anti-robot
MAX_ATTENTES_NOPECHA = 3       # attentes NopeCHA max par page avant d'abandonner le CAPTCHA

logger = configurer_logs(RACINE / "logs", "verify_eligibility")


class ArretVerification(Exception):
    """Arrêt propre : les offres non tranchées restent « À vérifier » pour le prochain lancement."""


class PageAntiRobot(Exception):
    """Le site affiche une vérification anti-robot à la place de la page (Cloudflare…), que
    NopeCHA et le clic de secours n'ont pas franchie : l'offre n'a pas pu être lue."""


# ------------------------------------------------------------------------------------------------
# Instance unique
# ------------------------------------------------------------------------------------------------

def prendre_verrou() -> bool:
    for _ in range(2):
        try:
            fd = os.open(str(VERROU), os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.write(fd, f"verify_eligibility {os.getpid()} {datetime.now():%Y-%m-%d %H:%M:%S}".encode())
            os.close(fd)
            return True
        except FileExistsError:
            try:
                morceaux = VERROU.read_text(encoding="utf-8", errors="replace").split()
                vivant = len(morceaux) >= 2 and morceaux[1].isdigit() and S._processus_vivant(int(morceaux[1]))
            except OSError:
                vivant = True
            if vivant:
                return False
            VERROU.unlink(missing_ok=True)       # verrou d'un lancement arrêté brutalement
    return False


# ------------------------------------------------------------------------------------------------
# Offres à vérifier
# ------------------------------------------------------------------------------------------------

def excel() -> em.ExcelManager:
    config, _ = charger_config()
    return em.ExcelManager(chemin_projet(config.get("chemins", {}).get("excel", "input/entreprises.xlsx")),
                           logger)


def offres_a_verifier() -> list[dict]:
    """Lignes « À vérifier », dans l'ordre de l'Excel, avec ce que le scraping a mémorisé."""
    try:
        memo = json.loads(S.MEMO_OFFRES.read_text(encoding="utf-8")) if S.MEMO_OFFRES.exists() else {}
    except (OSError, ValueError):
        memo = {}
    offres = []
    for d in excel().lire_toutes():
        if em.statut_normalise(d["Statut"]) != em.A_VERIFIER or not str(d["Lien_Offre"]).strip():
            continue
        cle = S.identifiant(str(d["Lien_Offre"]))
        m = memo.get(cle, {})
        commentaire = str(d.get("Commentaire") or "")
        titre = m.get("titre") or commentaire.rsplit(" — ", 1)[-1].strip()
        offres.append({"cle": cle, "lien": str(d["Lien_Offre"]).strip(), "titre": titre,
                       "entreprise": m.get("entreprise") or str(d["Nom_Entreprise"]).strip(),
                       "easy_apply": bool(m.get("easy_apply")) or "easy apply" in commentaire.lower(),
                       "secteur": m.get("secteur", ""), "description": m.get("description", ""),
                       "verdict": "", "raison": "", "etape": ""})
    for n, o in enumerate(offres, start=1):
        o["n"] = n
    return offres


# ------------------------------------------------------------------------------------------------
# DeepSeek
# ------------------------------------------------------------------------------------------------

def consigne() -> str:
    c = json.loads(S.CRITERES.read_text(encoding="utf-8"))
    lignes = [("Domaine souhaité", c.get("domaine_souhaite")), ("À inclure aussi", c.get("inclure")),
              ("À exclure", c.get("exclure")), ("Postes mixtes", c.get("cas_mixtes")),
              ("Précisions", c.get("precisions"))]
    criteres = "\n".join(f"- {nom} : {val}" for nom, val in lignes if str(val or "").strip())
    diplome = str(c.get("diplome") or "").strip()
    formation = (f"FORMATION DU CANDIDAT : {diplome}\n(Elle t'aide à comprendre son profil et le "
                 "vocabulaire de son domaine ; le verdict, lui, dépend des critères ci-dessous.)\n\n"
                 if diplome else "")
    return (
        f"Tu tries des offres {objectif_courant().de(objectif_courant().type_mot)} pour un candidat, "
        "selon SES critères de domaine.\n\n"
        + formation
        + f"CRITÈRES DU CANDIDAT :\n{criteres}\n\n"
        "VERDICTS :\n"
        "- \"true\" : l'offre correspond au domaine souhaité (y compris combinée avec un autre "
        "domaine, si les critères l'acceptent).\n"
        "- \"false\" : l'offre ne correspond pas, ou relève de ce qui est exclu.\n"
        "- \"pending\" : UNIQUEMENT si le texte fourni ne permet pas de savoir (ex. titre vague : "
        "« Stagiaire H/F », « Stage 6 mois », « Internship »). Ne devine jamais.\n\n"
        "RÈGLES : réponds pour CHAQUE numéro reçu, dans le MÊME ordre, sans en oublier ni en "
        "ajouter. Juge chaque offre seule, uniquement d'après ce qui est écrit.\n"
        "FORMAT : UN SEUL bloc ```json {\"resultats\": [{\"n\": 1, \"verdict\": \"true\", "
        "\"raison\": \"5 à 12 mots\"}]} et rien d'autre.")


def lire_verdicts(texte: str, attendus: list[int]) -> tuple[dict, str]:
    """({n: (verdict, raison)}, problème) — problème vide si la réponse est complète et valide."""
    try:
        donnees = extraire_json(texte)
    except ValueError as e:
        return {}, f"JSON illisible ({e})"
    liste = donnees.get("resultats") if isinstance(donnees, dict) else donnees
    if not isinstance(liste, list):
        return {}, "la clé « resultats » (une liste) est absente"
    res, invalides = {}, []
    for item in liste:
        if not isinstance(item, dict):
            continue
        try:
            n = int(item.get("n"))
        except (TypeError, ValueError):
            continue
        verdict = VERDICTS.get(str(item.get("verdict", "")).strip().lower())
        if verdict is None:
            invalides.append(n)
            continue
        res[n] = (verdict, " ".join(str(item.get("raison", "")).split())[:200])
    manquants = [n for n in attendus if n not in res]
    en_trop = sorted(set(res) - set(attendus))
    problemes = []
    if manquants:
        problemes.append(f"numéros manquants : {manquants}")
    if en_trop:
        problemes.append(f"numéros qui n'existent pas dans ce message : {en_trop}")
    if invalides:
        problemes.append(f"verdict invalide (true / false / pending) pour : {invalides}")
    return {n: v for n, v in res.items() if n in attendus}, " ; ".join(problemes)


def demander(chat, message: str, attendus: list[int]) -> dict:
    """Envoie le message ; redemande tant que la réponse ne couvre pas EXACTEMENT les numéros."""
    systeme = consigne()
    for essai in range(1, ESSAIS + 1):
        try:
            texte = chat.generer(systeme, message, images=None, tache="formulaire")
        except ErreurGeminiChat as e:
            raise ArretVerification(f"DeepSeek indisponible ({e})") from e
        res, probleme = lire_verdicts(texte, attendus)
        if not probleme:
            return res
        logger.warning(f"  Réponse de DeepSeek incomplète (essai {essai}/{ESSAIS}) : {probleme}.")
        message = (f"Ta réponse est invalide : {probleme}. Renvoie la liste COMPLÈTE, un verdict pour "
                   f"CHACUN des numéros {attendus}, dans cet ordre, au même format JSON.")
    raise ArretVerification(f"réponse de DeepSeek toujours incomplète après {ESSAIS} essais")


def appliquer(offres: list[dict], res: dict, etape: str) -> None:
    for o in offres:
        if o["n"] in res:
            o["verdict"], o["raison"] = res[o["n"]]
            o["etape"] = etape


# ------------------------------------------------------------------------------------------------
# Étapes
# ------------------------------------------------------------------------------------------------

def etape_titres(chat, offres: list[dict]) -> None:
    lots = [offres[i:i + LOT_TITRES] for i in range(0, len(offres), LOT_TITRES)]
    for k, lot in enumerate(lots, start=1):
        lignes = "\n".join(f"{o['n']}. {o['titre']} — {o['entreprise']}" for o in lot)
        message = (f"ÉTAPE 1 — TITRES (lot {k}/{len(lots)}, numéros {lot[0]['n']} à {lot[-1]['n']}). "
                   f"Pour chaque offre : titre — entreprise.\n\n{lignes}")
        logger.info(f"Étape 1 : titres {lot[0]['n']} à {lot[-1]['n']} envoyés à DeepSeek…")
        appliquer(lot, demander(chat, message, [o["n"] for o in lot]), "titre")


def a_relire(o: dict) -> bool:
    """
    Offre dont la DESCRIPTION doit être lue : titre indéterminé (pending), ou écartée sur son seul
    titre — un titre peut tromper (« Études opérationnelles » pour un poste de Machine Learning) :
    aucune offre n'est écartée sans que sa description, quand on l'a, ait été lue.
    """
    return ((o["verdict"] == "pending" or (o["verdict"] == "false" and o["etape"] == "titre"))
            and len(o["description"]) > 100)


def etape_descriptions(chat, offres: list[dict]) -> None:
    cibles = [o for o in offres if a_relire(o)]
    for i in range(0, len(cibles), LOT_DESCRIPTIONS):
        lot = cibles[i:i + LOT_DESCRIPTIONS]
        blocs = "\n\n".join(
            f"### {o['n']}. {o['titre']} — {o['entreprise']}"
            + (f" (secteur : {o['secteur']})" if o["secteur"] else "")
            + f"\n{o['description'][:3500]}" for o in lot)
        message = ("ÉTAPE 2 — DESCRIPTIONS : le titre de ces offres ne suffisait pas, ou pouvait "
                   "tromper. Voici leur description, telle que publiée sur LinkedIn. Juge chaque "
                   "offre d'après SA description, c'est-à-dire d'après les missions et les "
                   "compétences réellement demandées, même si le titre ou le service semble "
                   "éloigné du domaine (pending seulement si elle ne dit toujours rien du poste)."
                   f"\n\n{blocs}")
        logger.info(f"Étape 2 : descriptions des offres {[o['n'] for o in lot]}…")
        appliquer(lot, demander(chat, message, [o["n"] for o in lot]), "description")


JS_LIEN_SITE = r"""
(() => {
  const liens = [...document.querySelectorAll('a[href]')];
  const offsite = liens.find(a => /offsite|externalApply/i.test(
      (a.getAttribute('data-tracking-control-name') || '') + ' ' + a.href));
  if (offsite) return offsite.href;
  const postuler = liens.find(a => /postuler|apply/i.test(a.innerText || '')
                                   && !/linkedin\.com/i.test(a.href) && /^https?:/.test(a.href));
  return postuler ? postuler.href : '';
})()
"""


def passer_captcha(nav: Navigateur, n: int) -> None:
    """
    CAPTCHA sur la page qui vient d'être ouverte : EXACTEMENT la logique du robot de candidature.
      1) Page anti-robot À LA PLACE du contenu (orchestrator.extraire_offre, lecture de l'offre) :
         attente de NopeCHA, puis UN clic sur la case en secours ; toujours là -> PageAntiRobot.
      2) CAPTCHA affiché dans la page (candidature.postuler, avant chaque étape) : jusqu'à
         MAX_ATTENTES_NOPECHA attentes de NopeCHA, puis UN clic sur la case en secours.
    """
    # 1) Comme orchestrator.py (lecture de l'offre).
    if nav.page_anti_robot():
        # NopeCHA (si actif) résout seule la vérification ; le clic sur la case reste en secours.
        clic = ("NopeCHA " + nav.attendre_nopecha()) if nav.nopecha_actif else ""
        if clic == "NopeCHA validée":
            nav.attendre_page_prete(15)
        if nav.page_anti_robot():
            clic = (clic + " ; " if clic else "") + nav.cliquer_case_anti_robot()
            if clic.endswith("case cliquée"):
                clic += " -> " + nav.attendre_case_validee(DELAI_CASE_VALIDEE)
                if clic.endswith("-> validée"):
                    nav.attendre_page_prete(15)   # la page se recharge
        logger.info(f"Étape 3 : offre {n} — page de vérification anti-robot -> {clic}.")
        if nav.page_anti_robot():
            raise PageAntiRobot(f"Page de vérification anti-robot du site ({clic}) : "
                                f"offre non lue, à ouvrir à la main.")
    # 2) Comme candidature.py (CAPTCHA visible avant chaque étape du formulaire).
    try:
        captcha = nav.captcha_a_resoudre()
    except Exception:  # noqa: BLE001
        captcha = False
    attentes = 0
    while captcha and nav.nopecha_actif and attentes < MAX_ATTENTES_NOPECHA:
        attentes += 1
        nopecha = nav.attendre_nopecha()
        if nopecha in ("validée", "disparue"):
            logger.info(f"Étape 3 : offre {n} — CAPTCHA visible -> résolu par NopeCHA ({nopecha}).")
            captcha = False
        else:
            logger.info(f"Étape 3 : offre {n} — NopeCHA : CAPTCHA {nopecha} "
                        f"({attentes}/{MAX_ATTENTES_NOPECHA}).")
    # Secours : UN clic sur la case, une fois par page.
    if captcha:
        clic = nav.cliquer_case_anti_robot()
        if clic == "case cliquée":
            clic += " -> " + nav.attendre_case_validee(DELAI_CASE_VALIDEE)
        logger.info(f"Étape 3 : offre {n} — case anti-robot visible -> {clic} "
                    f"(une seule fois sur cette page).")


def lien_site_entreprise(nav: Navigateur, lien_linkedin: str, n: int) -> str:
    """Lien « Postuler sur le site de l'entreprise » de la page PUBLIQUE LinkedIn (sans connexion)."""
    nav.aller(lien_linkedin)
    nav.attendre_page(4.0)
    passer_captcha(nav, n)
    brut = str(nav.b.eval(JS_LIEN_SITE) or "")
    if not brut.startswith("http"):
        return ""
    cible = parse_qs(urlsplit(brut).query).get("url", [""])[0]   # externalApply?url=<site>
    return cible if cible.startswith("http") else ("" if "linkedin.com" in brut else brut)


def etape_site(chat, nav: Navigateur, offres: list[dict]) -> None:
    for o in [o for o in offres if o["verdict"] == "pending" and not o["easy_apply"]]:
        try:
            site = lien_site_entreprise(nav, o["lien"], o["n"])
            if not site:
                logger.info(f"Étape 3 : offre {o['n']} — pas de lien vers le site de l'entreprise.")
                continue
            nav.aller(site)
            nav.attendre_page(5.0)
            passer_captcha(nav, o["n"])
            texte = (nav.texte_complet() or "").strip()[:9000]
        except PageAntiRobot as e:
            # Comme le robot de candidature : page non lue -> « Action manuelle », rien n'est
            # envoyé à DeepSeek.
            o["verdict"], o["raison"], o["etape"] = "manuel", str(e), "site de l'entreprise"
            logger.info(f"Étape 3 : offre {o['n']} -> {em.ACTION_MANUELLE} : {e}")
            continue
        except Exception as e:  # noqa: BLE001 — une page qui résiste ne bloque pas le tri
            logger.info(f"Étape 3 : offre {o['n']} — site illisible ({type(e).__name__}).")
            continue
        if len(texte) < 200:
            continue
        message = (f"ÉTAPE 3 — SITE DE L'ENTREPRISE pour l'offre {o['n']} : « {o['titre']} » chez "
                   f"{o['entreprise']}. Texte de la page du site ({site[:120]}) ci-dessous.\n"
                   "ATTENTION : cette page peut présenter PLUSIEURS postes. Ne juge QUE le poste "
                   f"« {o['titre']} » ; si tu ne le trouves pas sur la page, verdict pending.\n\n{texte}")
        logger.info(f"Étape 3 : offre {o['n']} — page de l'entreprise envoyée à DeepSeek…")
        appliquer([o], demander(chat, message, [o["n"]]), "site de l'entreprise")


# ------------------------------------------------------------------------------------------------
# Écriture dans l'Excel
# ------------------------------------------------------------------------------------------------

def enregistrer(offres: list[dict]) -> tuple[int, int]:
    """
    true -> « À traiter » ; false -> feuille « Écartées » (ligne retirée de « Offres »).
    Chaque ligne est retrouvée par son LIEN, sous le verrou commun avec le robot et le scraping,
    et seulement si elle est TOUJOURS « À vérifier » (jamais une ligne modifiée entre-temps).
    """
    x = excel()
    gardees = ecartees = manuelles = 0
    with x._verrou():
        wb, ws, entetes = x._ouvrir()
        par_cle = {}
        for n in range(2, ws.max_row + 1):
            lien = ws.cell(row=n, column=entetes["Lien_Offre"]).value
            if lien and em.statut_normalise(ws.cell(row=n, column=entetes["Statut"]).value) == em.A_VERIFIER:
                par_cle.setdefault(S.identifiant(str(lien)), n)
        colonnes = [nom for nom, _ in sorted(entetes.items(), key=lambda kv: kv[1])]
        if S.FEUILLE_ECARTEES in wb.sheetnames:
            ec = wb[S.FEUILLE_ECARTEES]
        else:
            ec = wb.create_sheet(S.FEUILLE_ECARTEES)
            for j, nom in enumerate(colonnes + ["Raison_Ecart", "Date_Ecart"], start=1):
                ec.cell(row=1, column=j, value=nom).font = Font(bold=True)
            ec.freeze_panes = "A2"
        col_ec = {str(c.value).strip(): c.column for c in ec[1] if c.value}
        maintenant = datetime.now().strftime("%Y-%m-%d %H:%M")
        a_retirer = []
        for o in offres:
            n = par_cle.get(o["cle"])
            if n is None or o["verdict"] not in ("true", "false", "manuel"):
                continue
            raison = o["raison"] or ("hors domaine" if o["verdict"] == "false" else "dans le domaine")
            if o["verdict"] == "manuel":
                # Page anti-robot non franchie (étape 3) : comme le robot de candidature,
                # « Action manuelle » — ni gardée, ni écartée : l'utilisateur décide.
                ancien = str(ws.cell(row=n, column=entetes["Commentaire"]).value or "")
                ws.cell(row=n, column=entetes["Statut"], value=em.ACTION_MANUELLE)
                ws.cell(row=n, column=entetes["Commentaire"],
                        value=(ancien + f" — Domaine non vérifié ({o['etape']}) : {raison}")[:32000])
                if "Date_Traitement" in entetes:
                    ws.cell(row=n, column=entetes["Date_Traitement"], value=maintenant)
                manuelles += 1
            elif o["verdict"] == "true":
                ancien = str(ws.cell(row=n, column=entetes["Commentaire"]).value or "")
                ws.cell(row=n, column=entetes["Statut"], value=em.A_TRAITER)
                ws.cell(row=n, column=entetes["Commentaire"],
                        value=(ancien + f" — Domaine vérifié ({o['etape']}) : {raison}")[:32000])
                gardees += 1
            else:
                dest = ec.max_row + 1
                for nom in colonnes:
                    if nom in col_ec:
                        ec.cell(row=dest, column=col_ec[nom], value=ws.cell(row=n, column=entetes[nom]).value)
                ec.cell(row=dest, column=col_ec["Statut"], value="Écartée")
                ec.cell(row=dest, column=col_ec["Raison_Ecart"], value=f"{raison} (étape : {o['etape']})")
                ec.cell(row=dest, column=col_ec["Date_Ecart"], value=maintenant)
                a_retirer.append(n)
                ecartees += 1
        for n in sorted(a_retirer, reverse=True):        # du bas vers le haut : aucun décalage
            ws.delete_rows(n, 1)
        # « À vérifier » dans la liste déroulante du Statut (sinon Excel la signale invalide)
        for dv in ws.data_validations.dataValidation:
            f = str(dv.formula1 or "")
            if em.A_TRAITER in f and em.A_VERIFIER not in f and len(f) + 12 < 255:
                dv.formula1 = f[:-1] + "," + em.A_VERIFIER + '"'
        x._enregistrer(wb)
    if manuelles:
        logger.info(f"{manuelles} offre(s) en « {em.ACTION_MANUELLE} » (page anti-robot non "
                    f"franchie : à ouvrir à la main).")
    return gardees, ecartees


# ------------------------------------------------------------------------------------------------
# Programme
# ------------------------------------------------------------------------------------------------

def ouvrir_navigateur() -> tuple[Navigateur, dict]:
    config, _ = charger_config()
    config = copy.deepcopy(config)
    config.setdefault("navigateur", {}).update({
        "cdp_url": f"http://127.0.0.1:{PORT}", "profil_chrome": PROFIL,
        "navigateurs_secours": [], "lancer_automatiquement": True, "headless": False})
    # NopeCHA : réglage de config.yaml (section nopecha), comme le robot de candidature.
    nav = Navigateur(config, logger)
    nav.demarrer(premier_plan=False)                         # ne vole pas l'écran au robot
    return nav, config


def fermer_navigateur(nav: Navigateur) -> None:
    try:
        cdp = nav.b._connect_browser()
        try:
            cdp.send("Browser.close")
        finally:
            cdp.close()
    except Exception:  # noqa: BLE001
        pass


def trier(chat, nav: Navigateur, offres: list[dict]) -> tuple[int, int]:
    logger.info(f"{len(offres)} offre(s) « {em.A_VERIFIER} » à trier.")
    etape_titres(chat, offres)
    n_pending = sum(o["verdict"] == "pending" for o in offres)
    logger.info(f"Étape 1 terminée : {sum(o['verdict'] == 'true' for o in offres)} true, "
                f"{sum(o['verdict'] == 'false' for o in offres)} false, {n_pending} pending.")
    n_relus = sum(o["verdict"] == "false" and a_relire(o) for o in offres)
    if n_relus:
        logger.info(f"{n_relus} offre(s) écartée(s) sur le titre : description relue avant de "
                    "l'écarter (un titre peut tromper).")
    if n_pending or n_relus:
        etape_descriptions(chat, offres)
        etape_site(chat, nav, offres)
    for o in offres:
        if o["verdict"] == "pending":
            o["verdict"], o["raison"], o["etape"] = "false", RAISON_INCONNU, o["etape"] or "titre"
    for o in offres:
        logger.info(f"  {o['n']:>3}. [{o['verdict'].upper():5}] {o['titre'][:60]} — {o['entreprise'][:30]}"
                    f" | {o['raison'][:70]}")
    return enregistrer(offres)


def main() -> int:
    if not S.CRITERES.exists():
        logger.error("criteres_domaine.json absent : lance input/Scraping/PARAMETRER_SCRAPING.bat.")
        return 1
    if not prendre_verrou():
        logger.info("Une vérification est déjà en cours : rien à faire ici.")
        return 0
    nav = None
    code = 0
    offres: list[dict] = []
    try:
        offres = offres_a_verifier()
        if not offres:
            logger.info(f"Aucune offre « {em.A_VERIFIER} » : rien à trier.")
            return 0
        logger.info("===== Vérification du domaine des offres scrapées =====")
        nav, config = ouvrir_navigateur()
        chat = ClientDeepSeekFormulaire(config, nav, logger)
        total_g = total_e = 0
        for tour in range(1, TOURS_MAX + 1):
            offres = offres if tour == 1 else offres_a_verifier()   # arrivées pendant le tri
            if not offres:
                break
            g, e = trier(chat, nav, offres)
            total_g, total_e = total_g + g, total_e + e
        logger.info(f"===== Terminé : {total_g} offre(s) gardée(s) (« {em.A_TRAITER} »), "
                    f"{total_e} écartée(s) (feuille « {S.FEUILLE_ECARTEES} ») =====")
        try:
            chat.fermer()
        except Exception:  # noqa: BLE001
            pass
    except ArretVerification as e:
        code = 2
        # Ce qui est déjà tranché est gardé ; le reste (pending, lots non envoyés) reste à vérifier.
        tranchees = [o for o in offres if o["verdict"] in ("true", "false", "manuel")]
        if tranchees:
            try:
                g, ec = enregistrer(tranchees)
                logger.info(f"Déjà tranchées avant l'arrêt et enregistrées : {g} gardée(s), {ec} écartée(s).")
            except Exception as e2:  # noqa: BLE001
                logger.error(f"Enregistrement partiel impossible ({e2}).")
        logger.error(f"Vérification arrêtée : {e}. Les offres non triées restent « {em.A_VERIFIER} » "
                     f"et seront reprises au prochain lancement.")
        logger.error("Si DeepSeek n'est pas connecté dans ce navigateur : LANCER.bat -> 8 -> 5, "
                     "connecte-toi une fois à DeepSeek, puis LANCER.bat -> 12.")
    except Exception as e:  # noqa: BLE001 — jamais de plantage silencieux
        code = 3
        logger.exception(f"Vérification interrompue : {type(e).__name__} : {e}. Les offres non triées "
                         f"restent « {em.A_VERIFIER} ».")
    finally:
        if nav is not None and code == 0:
            fermer_navigateur(nav)        # en cas d'échec, la fenêtre reste ouverte (connexion…)
        VERROU.unlink(missing_ok=True)
    time.sleep(8 if code == 0 else 30)    # le temps de lire le bilan, puis la fenêtre se ferme
    return code


if __name__ == "__main__":
    sys.exit(main())

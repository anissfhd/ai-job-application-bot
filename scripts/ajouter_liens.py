"""
Ajoute des liens d'offres dans l'Excel, en COLLANT une liste dans la console.

Usage : lancé par l'entrée « Coller des liens d'offres » du menu LANCER.bat.
On colle le bloc de liens (un par ligne), on appuie sur Entrée, et c'est écrit.

Seule la colonne Lien_Offre est remplie : le nom de l'entreprise, les descriptions et le
statut restent vides — le robot les complétera lui-même à la lecture de l'offre. Un statut
vide vaut « À traiter », donc les liens ajoutés partent au prochain lancement.
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))

from utils import charger_config, chemin_projet  # noqa: E402
from excel_manager import enregistrer_sans_risque  # noqa: E402

COLONNE_LIEN = "Lien_Offre"


# Paramètres de suivi : ils changent d'un partage à l'autre sans désigner une autre offre.
SUIVI = ("utm_", "trk", "trackingid", "refid", "originalsubdomain", "position",
         "pagenum", "eboriginalsubdomain", "savedsearchid", "fbclid", "gclid")


def nettoyer(url: str) -> str:
    """
    Ramène un lien d'offre à sa forme MINIMALE et durable.

    Les paramètres de suivi (trackingId, ctt, cb…) sont horodatés : ils périment, changent
    d'un partage à l'autre, et font passer deux fois la même offre pour deux offres
    différentes. On ne garde donc que ce qui identifie réellement l'offre — et cet endroit
    DÉPEND DU SITE :
        LinkedIn  : dans le chemin       -> /jobs/view/<id>/
        Indeed    : dans la requête      -> ?jk=<id>
        Glassdoor : dans la requête      -> ?jl=<id>
    Couper après « ? » partout, comme on pourrait le croire, détruirait les deux derniers.

    Site inconnu : on retire seulement les paramètres de suivi connus, sans rien décoder —
    ré-encoder une URL qu'on ne comprend pas est le meilleur moyen de la casser.
    """
    url = url.strip()
    sans_protocole = url.split("://", 1)[-1]
    hote = sans_protocole.split("/", 1)[0].lower()
    reste = sans_protocole[len(hote):]
    chemin, _, requete = reste.partition("?")
    parts = [p for p in requete.split("&") if p]

    def parametre(*noms: str) -> str:
        for nom in noms:
            for p in parts:
                if p.lower().startswith(nom.lower() + "="):
                    return p
        return ""

    if "linkedin.com" in hote:
        segments = [s for s in chemin.split("/") if s]
        if len(segments) >= 3 and segments[0] == "jobs" and segments[1] == "view":
            # Lien public « /jobs/view/<titre>-<id> » : seul le numéro final identifie l'offre.
            # (Même règle dans input/Scraping/scraping.py — à garder identique.)
            numero = segments[2].rsplit("-", 1)[-1] if not segments[2].isdigit() else segments[2]
            return f"https://www.linkedin.com/jobs/view/{numero}/"
    elif "indeed." in hote:
        jk = parametre("jk")
        if jk:
            return f"https://{hote}{chemin}?{jk}"
    elif "glassdoor." in hote:
        jl = parametre("jobListingId", "jl")
        if jl:
            return f"https://{hote}{chemin}?{jl}"

    utiles = [p for p in parts if not any(p.lower().startswith(s) for s in SUIVI)]
    return f"https://{hote}{chemin}" + ("?" + "&".join(utiles) if utiles else "")


def identifiant(url: str) -> str:
    """Clé de comparaison : la forme nettoyée, en minuscules, sans « www. » ni « / » final."""
    propre = nettoyer(url).lower()
    for prefixe in ("https://", "http://"):
        if propre.startswith(prefixe):
            propre = propre[len(prefixe):]
    if propre.startswith("www."):
        propre = propre[4:]
    chemin, _, requete = propre.partition("?")
    return chemin.rstrip("/") + ("?" + requete if requete else "")


def lire_liens_colles() -> list[str]:
    """
    Lit les liens collés jusqu'à une ligne VIDE (ou la fin de l'entrée).

    C'est le seul mode qui accepte un collage multi-lignes : `set /p` du .bat ne lirait
    que la première ligne, et les suivantes seraient exécutées comme des commandes.
    """
    print("Colle tes liens (un par ligne), puis appuie sur Entrée sur une ligne vide.")
    print("(Ctrl+C pour annuler)\n")
    liens = []
    while True:
        try:
            ligne = input().strip()
        except EOFError:
            break
        if not ligne:
            break
        if not ligne.lower().startswith(("http://", "https://")):
            print(f"   ignoré (ce n'est pas un lien) : {ligne[:70]}")
            continue
        liens.append(ligne)
    return liens


def main() -> int:
    config, _ = charger_config()
    xlsx = chemin_projet(config["chemins"]["excel"])
    if not xlsx.exists():
        print(f"Fichier introuvable : {xlsx}")
        return 1

    liens = lire_liens_colles()
    print(f"\n{len(liens)} lien(s) reçu(s)." if liens
          else "\nAucun lien collé — je vérifie quand même les doublons du fichier.")

    wb = openpyxl.load_workbook(xlsx)      # surtout pas data_only : on préserve le fichier
    ws = wb.active
    entetes = [c.value for c in ws[1]]
    if COLONNE_LIEN not in entetes:
        print(f"Colonne « {COLONNE_LIEN} » absente de la feuille.")
        return 1
    col = entetes.index(COLONNE_LIEN) + 1

    # Les liens déjà écrits sont remis à la forme propre : sans ça, un lien collé hier avec
    # ses paramètres de suivi et le même collé aujourd'hui sans eux resteraient deux lignes.
    nettoyes = 0
    for r in range(2, ws.max_row + 1):
        lien = ws.cell(row=r, column=col).value
        if lien:
            propre = nettoyer(str(lien))
            if propre != str(lien):
                ws.cell(row=r, column=col).value = propre
                nettoyes += 1
    if nettoyes:
        print(f"\n{nettoyes} lien(s) déjà présent(s) raccourci(s) à leur forme durable.")

    # Dernière ligne réellement remplie (une ligne seulement mise en forme ne compte pas)
    derniere, existants = 1, {}
    for r in range(2, ws.max_row + 1):
        valeurs = [ws.cell(row=r, column=c).value for c in range(1, len(entetes) + 1)]
        if any(v not in (None, "") for v in valeurs):
            derniere = r
        lien = ws.cell(row=r, column=col).value
        if lien:
            existants[identifiant(str(lien))] = r

    a_ecrire, doublons, vus = [], [], set()
    for url in liens:
        cle = identifiant(url)
        if cle in existants:
            doublons.append((url, f"déjà en ligne {existants[cle]}"))
        elif cle in vus:
            doublons.append((url, "en double dans ta liste"))
        else:
            vus.add(cle)
            a_ecrire.append(nettoyer(url))   # on enregistre la forme propre, pas le lien brut

    if doublons:
        print(f"\n{len(doublons)} doublon(s) écarté(s) — une candidature en double serait envoyée :")
        for url, ou in doublons:
            print(f"   {url}  ({ou})")

    a_supprimer = doublons_du_fichier(ws, entetes, col)
    if not a_ecrire and not a_supprimer and not nettoyes:
        print("\nRien à faire — aucun lien nouveau, aucun doublon, rien à raccourcir.")
        return 0

    ligne = derniere + 1
    for url in a_ecrire:
        ws.cell(row=ligne, column=col).value = url
        ligne += 1

    # Suppression APRÈS l'écriture, en partant du bas : supprimer une ligne décale
    # toutes celles du dessous, et les numéros relevés plus haut deviendraient faux.
    for r in sorted(a_supprimer, reverse=True):
        ws.delete_rows(r)

    try:
        enregistrer_sans_risque(wb, xlsx)
    except PermissionError:
        print(f"\nÉCHEC : « {xlsx.name} » est ouvert dans Excel. Ferme-le et recommence.")
        print("Rien n'a été modifié.")
        return 1

    if a_ecrire:
        print(f"\n{len(a_ecrire)} lien(s) ajouté(s) : lignes {derniere + 1} à {ligne - 1}.")
        print("Seule la colonne Lien_Offre a été remplie ; le statut vide vaut « À traiter ».")
    if a_supprimer:
        print(f"{len(a_supprimer)} ligne(s) en double supprimée(s).")
    if nettoyes:
        print(f"{nettoyes} lien(s) existant(s) raccourci(s).")
    return 0


# Une ligne déjà traitée vaut mieux qu'une ligne vide : on garde celle qui porte le plus
# d'information, pour ne pas perdre l'historique ni risquer de repostuler.
RANG_STATUT = {"Postulé": 0, "Prêt (test)": 1, "Action manuelle": 2, "Erreur": 3,
               "Offre expirée": 4, "Ignoré": 5, "Fait": 6, "En cours": 7}


def doublons_du_fichier(ws, entetes: list, col: int) -> list[int]:
    """
    Repère les lignes qui désignent DEUX FOIS la même offre et renvoie celles à supprimer.

    On garde la ligne la plus « avancée » (Postulé avant Erreur, et tout avant une ligne
    vide), à égalité la première. Rien n'est supprimé sans un « o » explicite : ces lignes
    peuvent contenir l'historique d'une candidature déjà envoyée.
    """
    i_statut = entetes.index("Statut") + 1 if "Statut" in entetes else None
    i_nom = entetes.index("Nom_Entreprise") + 1 if "Nom_Entreprise" in entetes else None
    groupes: dict[str, list[int]] = {}
    for r in range(2, ws.max_row + 1):
        lien = ws.cell(row=r, column=col).value
        if lien:
            groupes.setdefault(identifiant(str(lien)), []).append(r)
    groupes = {k: v for k, v in groupes.items() if len(v) > 1}
    if not groupes:
        return []

    def valeur(cellule) -> str:
        return str(cellule.value or "").strip()

    a_supprimer = []
    print(f"\n{len(groupes)} offre(s) en double DANS le fichier :")
    for rangees in groupes.values():
        classees = sorted(rangees, key=lambda r: (
            RANG_STATUT.get(valeur(ws.cell(row=r, column=i_statut)) if i_statut else "", 99), r))
        garde, autres = classees[0], classees[1:]
        for r in classees:
            nom = valeur(ws.cell(row=r, column=i_nom)) if i_nom else ""
            statut = valeur(ws.cell(row=r, column=i_statut)) if i_statut else ""
            marque = "GARDÉE  " if r == garde else "supprimée"
            print(f"   {marque} ligne {r:<4} {nom[:22]:<22} {statut[:16]:<16} "
                  f"{valeur(ws.cell(row=r, column=col))[:46]}")
        a_supprimer += autres

    reponse = input(f"\nSupprimer ces {len(a_supprimer)} ligne(s) en double ? (o/N) ").strip().lower()
    if reponse != "o":
        print("Doublons conservés.")
        return []
    return a_supprimer


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print("\nAnnulé — rien n'a été modifié.")
        sys.exit(1)

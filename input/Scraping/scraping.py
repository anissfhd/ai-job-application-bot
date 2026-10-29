"""
Scraping automatique des offres LinkedIn via Apify, ajoutées dans l'Excel du robot.

Lancé par LANCER_SCRAPING.bat : on choisit l'heure du passage (ex. 09:00), puis la fenêtre reste
ouverte et le passage se refait chaque jour à la même heure (intervalle réglable dans
config_scraping.json : 1440 min = 24 h).

À chaque passage :
  1. l'acteur Apify « curious_coder/linkedin-jobs-scraper » est lancé sur les liens de liens.txt
     (filtre « publiées depuis le dernier passage » et tri par date ajoutés automatiquement) ;
  2. les liens sont nettoyés et comparés comme le fait l'option 9 du menu (« Coller des liens ») :
     une offre déjà présente dans l'Excel est ignorée ;
  3. les nouvelles offres, de la plus récente à la plus ancienne, sont INSÉRÉES juste sous la
     ligne où travaille le robot (statut « En cours ») — ou tout en haut s'il ne travaille pas.
     Les offres déjà présentes descendent d'autant.
  4. VÉRIFICATION DU DOMAINE (si « verification_domaine » est actif) : les nouvelles offres
     arrivent au statut « À vérifier », que le robot de candidature ignore. Le robot séparé
     scripts/verify_eligibility.py est lancé dans sa propre fenêtre : il fait trier les offres
     par DeepSeek selon criteres_domaine.json (réglé par PARAMETRER_SCRAPING.bat), passe les
     bonnes en « À traiter » et range les autres dans la feuille « Écartées », puis se ferme.
     Titre, description et entreprise de chaque offre sont gardés dans offres_scrapees.json.

Ce programme ne touche qu'à l'Excel. Il peut tourner pendant que le robot travaille : les deux
utilisent le même verrou (fichier « <excel>.verrou »), et le robot retrouve sa ligne par son lien
si elle a été décalée.
"""
from __future__ import annotations

import json
import logging
import os
import re
import sys
import time
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import parse_qsl, quote, urlencode, urlsplit, urlunsplit
from urllib.request import Request, urlopen

from openpyxl import load_workbook

DOSSIER = Path(__file__).resolve().parent
API = "https://api.apify.com/v2"

try:
    sys.stdout.reconfigure(encoding="utf-8")  # accents corrects dans la console
except Exception:  # noqa: BLE001
    pass


# ----------------------------------------------------------------------------------------------
# Configuration
# ----------------------------------------------------------------------------------------------

CONFIG_DEFAUT = {
    "excel": "../entreprises.xlsx",
    "acteur": "curious_coder~linkedin-jobs-scraper",
    "intervalle_minutes": 1440,
    "offres_max_par_lien": 20,
    "cout_max_par_passage_usd": 0.20,
    "filtrer_depuis_dernier_passage": True,
    "statut_nouvelles_offres": "À traiter",
    "candidats_max": 100,
    "priorite_easy_apply": True,
    "verification_domaine": True,
}

# Vérification du domaine des offres (robot séparé scripts/verify_eligibility.py)
STATUT_A_VERIFIER = "À vérifier"          # ignoré par le robot de candidature
FEUILLE_ECARTEES = "Écartées"             # offres hors domaine : gardées pour ne pas revenir
MEMO_OFFRES = DOSSIER / "offres_scrapees.json"
CRITERES = DOSSIER / "criteres_domaine.json"
VERIFICATEUR = DOSSIER.parent.parent / "scripts" / "verify_eligibility.py"
MEMO_JOURS = 45                           # une offre plus vieille est oubliée du mémo


def charger_config() -> dict:
    fichier = DOSSIER / "config_scraping.json"
    cfg = dict(CONFIG_DEFAUT)
    if fichier.exists():
        cfg.update(json.loads(fichier.read_text(encoding="utf-8")))
    return cfg


def charger_cle() -> str:
    cle = os.environ.get("APIFY_TOKEN", "")
    fichier = DOSSIER / ".env"
    if not cle and fichier.exists():
        for ligne in fichier.read_text(encoding="utf-8").splitlines():
            nom, _, valeur = ligne.partition("=")
            if nom.strip() == "APIFY_TOKEN":
                cle = valeur.strip().strip('"').strip("'")
    return cle


def charger_cles() -> list[str]:
    """
    Clés Apify DANS L'ORDRE : APIFY_TOKEN, puis APIFY_TOKEN_2 … APIFY_TOKEN_10 (fichier .env ou
    variables d'environnement). Si une clé est refusée (plus de solde, expirée, invalide), le
    passage essaie la suivante.
    """
    valeurs: dict[str, str] = {}
    fichier = DOSSIER / ".env"
    if fichier.exists():
        for ligne in fichier.read_text(encoding="utf-8").splitlines():
            nom, _, valeur = ligne.partition("=")
            if nom.strip().startswith("APIFY_TOKEN") and valeur.strip():
                valeurs[nom.strip()] = valeur.strip().strip('"').strip("'")
    noms = ["APIFY_TOKEN"] + [f"APIFY_TOKEN_{i}" for i in range(2, 11)]
    return [c for c in (os.environ.get(n, "") or valeurs.get(n, "") for n in noms) if c]


# Refus liés à la CLÉ (solde épuisé, clé expirée ou invalide) : on passe à la clé suivante.
REFUS_CLE = re.compile(r"credit|usage|limit|balance|payment|insufficient|exceed|token|invalid|expired|"
                       r"unauthori[sz]ed|forbidden|plan", re.IGNORECASE)


def charger_liens() -> list[str]:
    fichier = DOSSIER / "liens.txt"
    if not fichier.exists():
        return []
    return [l.strip() for l in fichier.read_text(encoding="utf-8").splitlines()
            if l.strip() and not l.strip().startswith("#")]


def journal() -> logging.Logger:
    logger = logging.getLogger("scraping")
    if not logger.handlers:
        logger.setLevel(logging.INFO)
        (DOSSIER / "logs").mkdir(exist_ok=True)
        fmt = logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s", "%Y-%m-%d %H:%M:%S")
        for h in (logging.StreamHandler(sys.stdout),
                  logging.FileHandler(DOSSIER / "logs" / "scraping.log", encoding="utf-8")):
            h.setFormatter(fmt)
            logger.addHandler(h)
    return logger


# ----------------------------------------------------------------------------------------------
# Liens : même nettoyage et même comparaison que l'option 9 du menu (scripts/ajouter_liens.py).
# Copiés ici pour que ce dossier ne dépende que de l'Excel. À garder identiques.
# ----------------------------------------------------------------------------------------------

SUIVI = ("utm_", "trk", "trackingid", "refid", "originalsubdomain", "position",
         "pagenum", "eboriginalsubdomain", "savedsearchid", "fbclid", "gclid")


def nettoyer(url: str) -> str:
    url = url.strip()
    sans_protocole = url.split("://", 1)[-1]
    hote = sans_protocole.split("/", 1)[0].lower()
    reste = sans_protocole[len(hote):]
    chemin, _, requete = reste.partition("?")
    parts = [p for p in requete.split("&") if p]

    def parametre(*noms: str) -> str:
        for nom in noms:
            for p in parts:
                if p.lower().startswith(nom.lower() + "="):
                    return p
        return ""

    if "linkedin.com" in hote:
        segments = [s for s in chemin.split("/") if s]
        if len(segments) >= 3 and segments[0] == "jobs" and segments[1] == "view":
            # Lien public « /jobs/view/<titre>-<id> » : seul le numéro final identifie l'offre.
            numero = segments[2].rsplit("-", 1)[-1] if not segments[2].isdigit() else segments[2]
            return f"https://www.linkedin.com/jobs/view/{numero}/"
    elif "indeed." in hote:
        jk = parametre("jk")
        if jk:
            return f"https://{hote}{chemin}?{jk}"
    elif "glassdoor." in hote:
        jl = parametre("jobListingId", "jl")
        if jl:
            return f"https://{hote}{chemin}?{jl}"

    utiles = [p for p in parts if not any(p.lower().startswith(s) for s in SUIVI)]
    return f"https://{hote}{chemin}" + ("?" + "&".join(utiles) if utiles else "")


def identifiant(url: str) -> str:
    propre = nettoyer(url).lower()
    for prefixe in ("https://", "http://"):
        if propre.startswith(prefixe):
            propre = propre[len(prefixe):]
    if propre.startswith("www."):
        propre = propre[4:]
    chemin, _, requete = propre.partition("?")
    return chemin.rstrip("/") + ("?" + requete if requete else "")


def variantes_easy_apply(url: str) -> tuple[str, str]:
    """
    (recherche Easy Apply seulement, recherche toutes offres) pour un même lien. L'acteur public
    ne dit pas si une offre est Easy Apply : une offre trouvée par la 1re recherche l'est forcément.
    """
    parties = urlsplit(url)
    params = [(k, v) for k, v in parse_qsl(parties.query, keep_blank_values=True) if k != "f_AL"]
    return (urlunsplit(parties._replace(query=urlencode(params + [("f_AL", "true")]))),
            urlunsplit(parties._replace(query=urlencode(params))))


def preparer_recherche(url: str, secondes: int) -> str:
    """
    Ajoute « publiées depuis le dernier passage » (f_TPR=r<secondes> : r86400 = 24 h) et le tri
    par date (sortBy=DD). 0 = lien laissé tel quel.
    """
    if not secondes:
        return url
    parties = urlsplit(url)
    params = [(k, v) for k, v in parse_qsl(parties.query, keep_blank_values=True)
              if k not in ("f_TPR", "sortBy")]
    params += [("f_TPR", f"r{int(secondes)}"), ("sortBy", "DD")]
    return urlunsplit(parties._replace(query=urlencode(params)))


# ----------------------------------------------------------------------------------------------
# Verrou commun avec le robot (scripts/excel_manager.py : verrou_excel). À garder identique.
# ----------------------------------------------------------------------------------------------

VERROU_PERIME = 600
JOURNAL = "scraping"


def _processus_vivant(pid: int) -> bool:
    """Le programme qui a posé le verrou tourne-t-il encore ? (Dans le doute : oui.)"""
    if pid <= 0:
        return True
    if os.name == "nt":
        import ctypes
        k = ctypes.windll.kernel32
        h = k.OpenProcess(0x1000, False, pid)          # PROCESS_QUERY_LIMITED_INFORMATION
        if not h:
            return k.GetLastError() == 5               # accès refusé = il existe
        code = ctypes.c_ulong()
        ok = k.GetExitCodeProcess(h, ctypes.byref(code))
        k.CloseHandle(h)
        return not ok or code.value == 259             # 259 = STILL_ACTIVE
    try:
        os.kill(pid, 0)
        return True
    except OSError:
        return False


def _verrou_orphelin(fichier: Path) -> bool:
    """Verrou dont le programme n'existe plus (fenêtre fermée, Ctrl+C, plantage) ou trop vieux."""
    try:
        if time.time() - fichier.stat().st_mtime > VERROU_PERIME:
            return True
        morceaux = fichier.read_text(encoding="utf-8", errors="replace").split()
        return len(morceaux) >= 2 and morceaux[1].isdigit() and not _processus_vivant(int(morceaux[1]))
    except OSError:
        return False


@contextmanager
def verrou_excel(chemin: Path, attente_max: float = 900.0):
    fichier = Path(str(chemin) + ".verrou")
    debut = time.monotonic()
    signale = False
    while True:
        try:
            fd = os.open(str(fichier), os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.write(fd, f"scraping {os.getpid()} {datetime.now():%Y-%m-%d %H:%M:%S}".encode())
            os.close(fd)
            break
        except FileExistsError:
            if _verrou_orphelin(fichier):
                try:
                    fichier.unlink(missing_ok=True)
                except OSError:
                    pass
                continue
            if not signale and time.monotonic() - debut > 5:
                signale = True
                try:
                    qui = fichier.read_text(encoding="utf-8", errors="replace").strip()
                except OSError:
                    qui = "?"
                logging.getLogger(JOURNAL).warning(
                    f"Excel verrouillé par un autre programme ({qui}) — attente de sa fin…")
            if time.monotonic() - debut > attente_max:
                raise TimeoutError(f"Verrou {fichier.name} toujours pris après {attente_max:.0f} s.")
            time.sleep(0.3)
    try:
        yield
    finally:
        try:
            fichier.unlink(missing_ok=True)
        except OSError:
            pass


# ----------------------------------------------------------------------------------------------
# Apify
# ----------------------------------------------------------------------------------------------

def appel(methode: str, url: str, donnees: dict | None = None, delai: float = 60) -> object:
    corps = json.dumps(donnees).encode() if donnees is not None else None
    req = Request(url, data=corps, method=methode, headers={"Content-Type": "application/json"})
    with urlopen(req, timeout=delai) as r:
        return json.loads(r.read().decode("utf-8"))


def lancer_apify(cle: str, cfg: dict, urls: list[str], log: logging.Logger) -> list[dict]:
    """Lance l'acteur, attend la fin (15 min au plus) et renvoie les offres trouvées."""
    entree = {"urls": urls, "limitPerSource": int(cfg["offres_max_par_lien"]), "scrapeCompany": False}
    params = {"token": cle}
    if cfg.get("cout_max_par_passage_usd"):
        params["maxTotalChargeUsd"] = cfg["cout_max_par_passage_usd"]
    run = appel("POST", f"{API}/acts/{quote(cfg['acteur'], safe='~')}/runs?{urlencode(params)}", entree)["data"]
    log.info(f"Apify : exécution {run['id']} lancée ({len(urls)} lien(s)).")
    fin = time.monotonic() + 15 * 60
    while run.get("status") in ("READY", "RUNNING"):
        if time.monotonic() > fin:
            raise TimeoutError("l'exécution Apify dépasse 15 min")
        time.sleep(10)
        run = appel("GET", f"{API}/actor-runs/{run['id']}?token={quote(cle)}")["data"]
    cout = (run.get("usageTotalUsd") or run.get("stats", {}).get("computeUnits") or "?")
    items = appel("GET", f"{API}/datasets/{run['defaultDatasetId']}/items?clean=true&format=json"
                         f"&token={quote(cle)}", delai=120)
    if run.get("status") != "SUCCEEDED":
        # Arrêt en cours de route (plafond de coût atteint, délai dépassé…) : on GARDE les offres
        # déjà récupérées au lieu de tout perdre ; les recherches suivantes seront incomplètes.
        if not items:
            raise RuntimeError(f"exécution Apify terminée avec le statut {run.get('status')}")
        log.warning(f"Apify : exécution arrêtée avant la fin (statut {run.get('status')}, coût {cout} $) "
                    f"— {len(items)} offre(s) déjà récupérée(s) conservée(s). Les dernières recherches "
                    f"de liens.txt peuvent manquer : augmente cout_max_par_passage_usd ou baisse "
                    f"offres_max_par_lien.")
        return items
    log.info(f"Apify : exécution terminée (coût : {cout} $).")
    return items


def lien_offre(offre: dict) -> str:
    for champ in ("link", "jobUrl", "url", "jobLink"):
        if str(offre.get(champ) or "").startswith("http"):
            return str(offre[champ])
    if offre.get("id"):
        return f"https://www.linkedin.com/jobs/view/{offre['id']}/"
    return ""


def date_publication(offre: dict) -> datetime | None:
    """Date de publication, quel que soit le format renvoyé (horodatage ms, ISO, AAAA-MM-JJ)."""
    for champ in ("postedAtTimestamp", "postedAt", "publishedAt", "listedAt", "postedDate", "postedTime"):
        v = offre.get(champ)
        if v in (None, ""):
            continue
        try:
            if isinstance(v, (int, float)) or str(v).isdigit():
                n = float(v)
                return datetime.fromtimestamp(n / 1000 if n > 1e11 else n, tz=timezone.utc)
            d = datetime.fromisoformat(str(v).replace("Z", "+00:00"))
            return d if d.tzinfo else d.replace(tzinfo=timezone.utc)
        except (ValueError, OSError, OverflowError):
            continue
    return None


# ----------------------------------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------------------------------

def enregistrer_sans_risque(wb, chemin: Path) -> None:
    """
    Écrit un fichier temporaire complet puis le substitue d'un coup à l'Excel (os.replace est
    atomique) : un Ctrl+C en plein enregistrement ne peut plus laisser un Excel illisible.
    """
    tmp = chemin.with_name(chemin.stem + ".enregistrement.xlsx")
    try:
        wb.save(tmp)
        os.replace(tmp, chemin)
    except BaseException:
        try:
            tmp.unlink(missing_ok=True)
        except OSError:
            pass
        raise


def inserer_dans_excel(chemin: Path, offres: list[dict], statut: str, log: logging.Logger) -> int:
    """
    Insère les offres (déjà triées, la plus récente d'abord) juste sous la ligne « En cours »,
    ou en ligne 2 si le robot ne travaille sur aucune. Ignore celles déjà présentes.
    Renvoie le nombre d'offres ajoutées.
    """
    with verrou_excel(chemin):
        wb = load_workbook(chemin)
        ws = wb["Offres"] if "Offres" in wb.sheetnames else wb.worksheets[0]
        col = {str(c.value).strip(): c.column for c in ws[1] if c.value}
        for nom in ("Nom_Entreprise", "Lien_Offre", "Statut", "Commentaire"):
            if nom not in col:
                raise KeyError(f"Colonne « {nom} » absente de l'Excel")

        deja = {identifiant(str(ws.cell(n, col["Lien_Offre"]).value))
                for n in range(2, ws.max_row + 1) if ws.cell(n, col["Lien_Offre"]).value}
        # Offres ÉCARTÉES (hors domaine) : déjà vues aussi, sinon chaque scraping les ramènerait.
        if FEUILLE_ECARTEES in wb.sheetnames:
            ecartees = wb[FEUILLE_ECARTEES]
            col_e = {str(c.value).strip(): c.column for c in ecartees[1] if c.value}
            if "Lien_Offre" in col_e:
                deja |= {identifiant(str(ecartees.cell(n, col_e["Lien_Offre"]).value))
                         for n in range(2, ecartees.max_row + 1)
                         if ecartees.cell(n, col_e["Lien_Offre"]).value}
        nouvelles, vues = [], set()
        for o in offres:
            cle = identifiant(o["lien"])
            if cle in deja or cle in vues:
                continue
            vues.add(cle)
            nouvelles.append(o)
        if not nouvelles:
            wb.close()
            return 0

        en_cours = next((n for n in range(2, ws.max_row + 1)
                         if str(ws.cell(n, col["Statut"]).value or "").strip().lower() == "en cours"), None)
        position = en_cours + 1 if en_cours else 2
        ws.insert_rows(position, amount=len(nouvelles))
        for i, o in enumerate(nouvelles):
            n = position + i
            ws.cell(n, col["Nom_Entreprise"], o["entreprise"])
            ws.cell(n, col["Lien_Offre"], nettoyer(o["lien"]))
            ws.cell(n, col["Statut"], statut)
            publie = o["date"].astimezone().strftime("%Y-%m-%d %H:%M") if o["date"] else "date inconnue"
            ws.cell(n, col["Commentaire"],
                    f"Scraping Apify {datetime.now():%Y-%m-%d %H:%M} — publiée {publie}"
                    f"{' — Easy Apply' if o.get('easy_apply') else ''} — {o['titre']}"[:300])

        for essai in range(1, 21):
            try:
                enregistrer_sans_risque(wb, chemin)
                break
            except PermissionError:
                log.warning(f"Excel ouvert ailleurs (dans Excel ?) — essai {essai}/20, nouvel essai "
                            f"dans 15 s. Ferme le fichier.")
                time.sleep(15)
        else:
            raise PermissionError("Excel toujours verrouillé : offres NON ajoutées ce passage.")
        ou = f"sous la ligne {en_cours} (robot en cours)" if en_cours else "en haut (ligne 2)"
        log.info(f"Excel : {len(nouvelles)} offre(s) insérée(s) {ou}, lignes {position} à "
                 f"{position + len(nouvelles) - 1}, de la plus récente à la plus ancienne.")
        return len(nouvelles)


# ----------------------------------------------------------------------------------------------
# Vérification du domaine : mémo des offres + lancement du robot séparé
# ----------------------------------------------------------------------------------------------

def verification_active(cfg: dict) -> bool:
    return bool(cfg.get("verification_domaine", True)) and CRITERES.exists() and VERIFICATEUR.exists()


def memoriser_offres(offres: list[dict], log: logging.Logger) -> None:
    """
    Garde titre, entreprise, Easy Apply et DESCRIPTION de chaque offre (clé = identifiant du
    lien) : verify_eligibility.py y lit tout, sans rouvrir LinkedIn. Mémo purgé après MEMO_JOURS.
    """
    try:
        memo = json.loads(MEMO_OFFRES.read_text(encoding="utf-8")) if MEMO_OFFRES.exists() else {}
    except (OSError, ValueError):
        memo = {}
    limite = (datetime.now() - timedelta(days=MEMO_JOURS)).strftime("%Y-%m-%d")
    memo = {k: v for k, v in memo.items() if str(v.get("scrape_le", "")) >= limite}
    maintenant = datetime.now().strftime("%Y-%m-%d %H:%M")
    for o in offres:
        memo[identifiant(o["lien"])] = {
            "lien": nettoyer(o["lien"]), "titre": o["titre"], "entreprise": o["entreprise"],
            "easy_apply": bool(o.get("easy_apply")), "secteur": o.get("secteur", ""),
            "description": o.get("description", "")[:8000], "scrape_le": maintenant}
    tmp = MEMO_OFFRES.with_suffix(".tmp")
    try:
        tmp.write_text(json.dumps(memo, ensure_ascii=False, indent=1), encoding="utf-8")
        os.replace(tmp, MEMO_OFFRES)
    except OSError as e:
        log.warning(f"Mémo des offres non enregistré ({e}) : la vérification relira les pages.")


def offres_a_verifier(chemin: Path) -> int:
    """Nombre de lignes « À vérifier » dans l'Excel (lecture seule, sans verrou)."""
    try:
        wb = load_workbook(chemin, read_only=True)
        ws = wb["Offres"] if "Offres" in wb.sheetnames else wb.worksheets[0]
        lignes = ws.iter_rows(values_only=True)
        entetes = [str(c or "").strip() for c in next(lignes)]
        i = entetes.index("Statut")
        n = sum(1 for r in lignes if r and len(r) > i and str(r[i] or "").strip() == STATUT_A_VERIFIER)
        wb.close()
        return n
    except Exception:  # noqa: BLE001
        return 0


def lancer_verification(chemin: Path, log: logging.Logger) -> None:
    """
    Lance verify_eligibility.py dans SA PROPRE fenêtre, détachée : ni ce programme ni le robot
    de candidature ne l'attendent, et un Ctrl+C sur eux ne l'arrête pas. Il ne tourne jamais
    en double (verrou géré par lui-même) et se ferme seul à la fin.
    """
    n = offres_a_verifier(chemin)
    if not n:
        return
    import subprocess
    drapeaux = (0x00000010 | 0x00000200) if os.name == "nt" else 0   # nouvelle console + groupe à part
    try:
        subprocess.Popen([sys.executable, str(VERIFICATEUR)], cwd=str(VERIFICATEUR.parent.parent),
                         creationflags=drapeaux, close_fds=True)
        log.info(f"Vérification du domaine lancée (fenêtre séparée) : {n} offre(s) « {STATUT_A_VERIFIER} ».")
    except OSError as e:
        log.error(f"Vérification du domaine NON lancée ({e}) : les offres restent « {STATUT_A_VERIFIER} ». "
                  f"Lance-la avec LANCER.bat, entrée 12.")


# ----------------------------------------------------------------------------------------------
# Passage et boucle
# ----------------------------------------------------------------------------------------------

def passage(log: logging.Logger) -> None:
    cfg = charger_config()
    cles = charger_cles()
    liens = charger_liens()
    excel = (DOSSIER / cfg["excel"]).resolve()
    if not cles:
        log.error("Clé Apify absente : écris APIFY_TOKEN=... dans input/Scraping/.env")
        return
    if not liens:
        log.error("Aucun lien de recherche : ajoute-les dans input/Scraping/liens.txt (un par ligne).")
        return
    if not excel.exists():
        log.error(f"Excel introuvable : {excel}")
        return

    fenetre = int(float(cfg["intervalle_minutes"]) * 60) if cfg.get("filtrer_depuis_dernier_passage", True) else 0
    if cfg.get("priorite_easy_apply", True):
        liens = [v for u in liens for v in variantes_easy_apply(u)]   # 2 recherches par lien
    urls = [preparer_recherche(u, fenetre) for u in liens]
    brutes, refusees = None, set()
    for numero, cle in enumerate(cles, start=1):
        if cle in refusees:
            log.info(f"Apify : clé n° {numero} identique à une clé déjà refusée -> ignorée.")
            continue
        try:
            brutes = lancer_apify(cle, cfg, urls, log)
            if numero > 1:
                log.info(f"Apify : passage réalisé avec la clé de secours n° {numero}.")
            break
        except HTTPError as e:
            detail = e.read().decode("utf-8", "replace")[:300]
            if e.code in (401, 402, 403) or REFUS_CLE.search(detail):
                refusees.add(cle)
                log.warning(f"Apify : clé n° {numero} refusée ({e.code} : {detail[:150]}) -> clé suivante.")
                continue
            log.error(f"Apify a refusé la demande ({e.code}) : {detail}")
            return
        except RuntimeError as e:
            # Exécution arrêtée sans aucune offre (souvent : solde épuisé en cours de route).
            if REFUS_CLE.search(str(e)) or "FAILED" in str(e) or "ABORTED" in str(e):
                refusees.add(cle)
                log.warning(f"Apify : clé n° {numero} : {e} -> clé suivante.")
                continue
            log.error(f"Apify : {e}")
            return
        except (URLError, TimeoutError) as e:
            log.error(f"Apify : {e}")
            return
    if brutes is None:
        log.error(f"Apify : aucune des {len(cles)} clé(s) n'a fonctionné (solde épuisé, expirée ou "
                  f"invalide). Ajoute ou remplace une clé dans input/Scraping/.env.")
        return

    offres, trop_de_candidats = [], 0
    candidats_max = int(cfg.get("candidats_max") or 0)
    for rang, o in enumerate(brutes if isinstance(brutes, list) else []):
        lien = lien_offre(o)
        if not lien:
            continue
        # « Moins de N candidats » : filtre absent de la recherche LinkedIn publique, appliqué ici.
        # « Over 200 applicants » ou « 150 » -> nombre ; inconnu -> gardée.
        chiffres = re.findall(r"\d+", str(o.get("applicantsCount") or "").replace(",", ""))
        if candidats_max and chiffres and int(chiffres[0]) >= candidats_max:
            trop_de_candidats += 1
            continue
        offres.append({"lien": lien, "rang": rang, "date": date_publication(o),
                       "easy_apply": "f_al=true" in str(o.get("inputUrl") or "").lower(),
                       "entreprise": str(o.get("companyName") or o.get("company") or "").strip(),
                       "titre": str(o.get("title") or o.get("jobTitle") or "").strip(),
                       "secteur": str(o.get("industries") or "").strip(),
                       "description": re.sub(r"\n\s*\n+", "\n\n",
                                             str(o.get("descriptionText") or "")).strip()})
    # Easy Apply d'abord (si priorite_easy_apply) ; puis la plus récente d'abord ; à date égale
    # (ou inconnue), l'ordre de LinkedIn trié par date. Une offre trouvée par les deux recherches
    # garde sa version Easy Apply (placée avant, le 1er exemplaire est gardé à l'insertion).
    plus_ancien = datetime.min.replace(tzinfo=timezone.utc)
    priorite = bool(cfg.get("priorite_easy_apply", True))
    offres.sort(key=lambda o: (0 if (priorite and o["easy_apply"]) else 1,
                               -(o["date"] or plus_ancien).timestamp() if o["date"] else float("inf"),
                               o["rang"]))
    if priorite:
        vues, uniques = set(), []
        for o in offres:
            if identifiant(o["lien"]) not in vues:
                vues.add(identifiant(o["lien"]))
                uniques.append(o)
        offres = uniques
        log.info(f"Easy Apply : {sum(o['easy_apply'] for o in offres)} offre(s), placées en premier ; "
                 f"autres : {sum(not o['easy_apply'] for o in offres)}.")
    log.info(f"Apify : {len(brutes) if isinstance(brutes, list) else 0} résultat(s), {len(offres)} retenue(s)"
             + (f", {trop_de_candidats} écartée(s) (≥ {candidats_max} candidats)" if trop_de_candidats else "") + ".")
    verifier = verification_active(cfg)
    if not offres:
        if verifier:
            lancer_verification(excel, log)      # offres restées « À vérifier » (arrêt précédent)
        return
    # Avec la vérification du domaine, les nouvelles offres attendent le tri (« À vérifier ») :
    # le robot de candidature ne les prend qu'une fois passées en « À traiter ».
    statut = STATUT_A_VERIFIER if verifier else cfg["statut_nouvelles_offres"]
    if verifier:
        memoriser_offres(offres, log)            # AVANT l'Excel : le vérificateur y lit tout
    try:
        ajoutees = inserer_dans_excel(excel, offres, statut, log)
    except Exception as e:  # noqa: BLE001 — un passage raté ne doit pas arrêter la boucle
        log.error(f"Écriture Excel impossible : {e}")
        return
    if not ajoutees:
        log.info("Aucune nouvelle offre : toutes sont déjà dans l'Excel (ou écartées).")
    if verifier:
        lancer_verification(excel, log)


def demander_heure() -> datetime:
    while True:
        saisie = input("Heure du passage quotidien (HH:MM, ex. 09:00 ; Entrée = maintenant) : ").strip()
        maintenant = datetime.now()
        if not saisie:
            return maintenant
        try:
            h, m = (int(x) for x in saisie.replace("h", ":").split(":"))
            debut = maintenant.replace(hour=h, minute=m, second=0, microsecond=0)
        except ValueError:
            print("  Format attendu : 14:30")
            continue
        if debut < maintenant:
            debut += timedelta(days=1)   # heure déjà passée aujourd'hui -> demain
        return debut


def main() -> int:
    log = journal()
    cfg = charger_config()
    print("=" * 64)
    print("  SCRAPING AUTOMATIQUE DES OFFRES (Apify -> Excel)")
    print(f"  Un passage toutes les {float(cfg['intervalle_minutes']) / 60:g} h, à la même heure. "
          f"Laisse cette fenêtre ouverte.")
    print("  Ctrl+C pour arrêter.")
    print("=" * 64)
    prochain = demander_heure()
    intervalle = timedelta(minutes=float(cfg["intervalle_minutes"]))
    log.info(f"Premier passage : {prochain:%d/%m %H:%M}, puis toutes les "
             f"{float(cfg['intervalle_minutes']) / 60:g} h.")
    try:
        while True:
            attente = (prochain - datetime.now()).total_seconds()
            if attente > 0:
                time.sleep(min(attente, 30))
                continue
            log.info("===== Passage de scraping =====")
            try:
                passage(log)
            except Exception as e:  # noqa: BLE001
                log.error(f"Passage interrompu : {type(e).__name__} : {e}")
            while prochain <= datetime.now():
                prochain += intervalle
            log.info(f"Prochain passage : {prochain:%d/%m %H:%M}.")
    except KeyboardInterrupt:
        log.info("Arrêt demandé (Ctrl+C).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
